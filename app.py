import os
import uuid
from flask import Flask, request, render_template, send_file, flash, redirect, url_for
from openpyxl import Workbook, load_workbook
import pytesseract

from parser import (procesar_factura, limpiar_numero, extraer_cabecera,
                    limpiar_cantidad, limpiar_sku)
from ocr_utils import es_imagen, convertir_imagen_a_pdf_ocr

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
EXCEL_PATH = os.path.join(OUTPUT_DIR, "facturas.xlsx")
ULTIMA_PATH = os.path.join(OUTPUT_DIR, "ultima_factura.xlsx")

os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

app = Flask(__name__)
app.secret_key = "dev"  # solo para uso local

CABECERA_COLUMNAS = [
    "archivo", "tipo_comprobante", "numero_factura", "fecha_facturacion",
    "cuit_emisor", "razon_social_cliente", "cuit_cliente", "condicion_iva_cliente",
    "neto", "iva", "subtotal", "total", "cae", "vencimiento_cae",
]

ITEMS_COLUMNAS = [
    "archivo", "numero_factura", "sku", "descripcion", "cantidad",
    "precio_unitario", "descuento", "neto", "interno", "iva", "subtotal",
]


def _crear_libro_con_hojas() -> Workbook:
    """Arma un Workbook nuevo con las dos hojas (Cabeceras / Items) y
    sus encabezados de columna, sin ninguna fila de datos todavía. La
    usan tanto el excel consolidado (que después se recarga y se le va
    agregando fila por fila) como el de "última factura" (que se arma
    de cero en cada subida)."""
    wb = Workbook()
    hoja_cab = wb.active
    hoja_cab.title = "Cabeceras"
    hoja_cab.append(CABECERA_COLUMNAS)

    hoja_items = wb.create_sheet("Items")
    hoja_items.append(ITEMS_COLUMNAS)

    return wb


def asegurar_excel():
    """Crea el excel CONSOLIDADO (con las 2 hojas) si no existe."""
    if os.path.exists(EXCEL_PATH):
        return
    _crear_libro_con_hojas().save(EXCEL_PATH)


# Columnas de items que son montos y tienen que ir al Excel como NUMERO,
# no como texto: si van como texto, Excel no las suma y ademas quedan con
# separadores inconsistentes segun como las haya leido el OCR
# ("9,968.23" en una fila y "9.722.83" en otra).
ITEMS_COLUMNAS_MONTO = ("precio_unitario", "neto", "interno", "iva", "subtotal")


def _valor_item(columna: str, valor):
    """Normaliza el valor de una celda de item antes de escribirlo al
    Excel: montos y cantidades como numero, sku sin la basura que el OCR
    pego de la descripcion, y el resto tal cual."""
    if valor is None:
        return None
    if columna in ITEMS_COLUMNAS_MONTO:
        return limpiar_numero(valor)
    if columna == "cantidad":
        return limpiar_cantidad(valor)
    if columna == "sku":
        return limpiar_sku(valor)
    return valor


def _filas_de(nombre_archivo: str, resultado: dict):
    """Arma la fila de cabecera y las filas de ítems para una factura ya
    procesada, en el orden de CABECERA_COLUMNAS/ITEMS_COLUMNAS. La
    comparten agregar_factura_a_excel (consolidado) y
    guardar_ultima_factura (para no repetir la lógica de armado)."""
    cab = resultado["cabecera"]
    fila_cab = [nombre_archivo] + [cab.get(c) for c in CABECERA_COLUMNAS[1:]]

    numero_factura = cab.get("numero_factura")
    filas_items = [
        [nombre_archivo, numero_factura] + [_valor_item(c, item.get(c))
                                            for c in ITEMS_COLUMNAS[2:]]
        for item in resultado["items"]
    ]
    return fila_cab, filas_items


def agregar_factura_a_excel(nombre_archivo: str, resultado: dict):
    """Suma esta factura como una fila más al Excel CONSOLIDADO
    (output/facturas.xlsx): se acumula entre subidas, nunca se pisa."""
    asegurar_excel()
    wb = load_workbook(EXCEL_PATH)
    hoja_cab = wb["Cabeceras"]
    hoja_items = wb["Items"]

    fila_cab, filas_items = _filas_de(nombre_archivo, resultado)
    hoja_cab.append(fila_cab)
    for fila_item in filas_items:
        hoja_items.append(fila_item)

    wb.save(EXCEL_PATH)


def guardar_ultima_factura(nombre_archivo: str, resultado: dict):
    """Guarda un Excel APARTE (output/ultima_factura.xlsx) con solo la
    factura que se acaba de procesar. A diferencia del consolidado, este
    archivo se REEMPLAZA por completo en cada subida — sirve para bajar
    rápido el resultado de la última factura sin tener que filtrar el
    consolidado ni reiniciarlo."""
    wb = _crear_libro_con_hojas()
    hoja_cab = wb["Cabeceras"]
    hoja_items = wb["Items"]

    fila_cab, filas_items = _filas_de(nombre_archivo, resultado)
    hoja_cab.append(fila_cab)
    for fila_item in filas_items:
        hoja_items.append(fila_item)

    wb.save(ULTIMA_PATH)


@app.route("/", methods=["GET"])
def index():
    return render_template(
        "index.html",
        excel_existe=os.path.exists(EXCEL_PATH),
        ultima_existe=os.path.exists(ULTIMA_PATH),
    )


@app.route("/subir", methods=["POST"])
def subir():
    archivo = request.files.get("factura")
    if not archivo or archivo.filename == "":
        flash("No seleccionaste ningún archivo.")
        return redirect(url_for("index"))

    nombre_lower = archivo.filename.lower()
    es_pdf = nombre_lower.endswith(".pdf")

    if not es_pdf and not es_imagen(nombre_lower):
        flash("El archivo debe ser un PDF o una foto/imagen (jpg, jpeg, png, webp, bmp, tif).")
        return redirect(url_for("index"))

    nombre_unico = f"{uuid.uuid4().hex[:8]}_{archivo.filename}"
    ruta_original = os.path.join(UPLOAD_DIR, nombre_unico)
    archivo.save(ruta_original)

    ruta_texto_debug = None
    ruta_items_debug = None
    items_ocr = []
    totales_ocr = {}
    texto_ocr_combinado = None

    try:
        if es_pdf:
            ruta_pdf = ruta_original
        else:
            # Es una foto/imagen: primero la convertimos en un PDF con OCR
            # (imagen + texto reconocido superpuesto) y de ahí en más se
            # procesa la CABECERA exactamente igual que cualquier otro PDF.
            # Los ÍTEMS y los TOTALES (Subtotal/IVA/Total), en cambio, se
            # reconstruyen aparte a partir de la posición de cada palabra
            # reconocida (ver ocr_items.py / ocr_totales.py), porque una
            # foto no tiene líneas de tabla reales.
            nombre_base = os.path.splitext(nombre_unico)[0]
            ruta_pdf = os.path.join(UPLOAD_DIR, f"{nombre_base}_ocr.pdf")
            ruta_texto_debug = os.path.join(UPLOAD_DIR, f"{nombre_base}_texto_ocr.txt")
            ruta_items_debug = os.path.join(UPLOAD_DIR, f"{nombre_base}_items_debug.json")
            (ruta_pdf, items_ocr, totales_ocr, advertencia_ocr,
             texto_ocr_combinado) = convertir_imagen_a_pdf_ocr(
                ruta_original, ruta_pdf, ruta_texto_debug
            )
            if advertencia_ocr:
                flash(advertencia_ocr)

        resultado = procesar_factura(ruta_pdf)

        if texto_ocr_combinado:
            # El texto combinado TIENE PRIORIDAD sobre el del PDF: el PDF
            # lleva solo la capa de Tesseract sobre la version gris, que
            # es la lectura MENOS confiable de las tres. Si se usa como
            # base y el texto combinado solo rellena huecos, un valor
            # equivocado pero no vacio de Tesseract nunca se corrige.
            # Caso real: Tesseract leia numero_factura = "021-083431",
            # que es el numero de IIBB, y tapaba el correcto
            # ("0011-00016528") que si habia leido PaddleOCR.
            cabecera_pdf = resultado["cabecera"]
            resultado["cabecera"] = extraer_cabecera(texto_ocr_combinado)
            for campo, valor in cabecera_pdf.items():
                if resultado["cabecera"].get(campo) is None and valor is not None:
                    resultado["cabecera"][campo] = valor

        if not resultado["items"] and items_ocr:
            resultado["items"] = items_ocr

        for campo, valor in totales_ocr.items():
            if resultado["cabecera"].get(campo) is None:
                resultado["cabecera"][campo] = limpiar_numero(valor)

        agregar_factura_a_excel(archivo.filename, resultado)
        guardar_ultima_factura(archivo.filename, resultado)

        # Solo avisamos de campos que realmente se escriben en el Excel
        # (algunas claves de `cabecera` son internas/auxiliares, como
        # condicion_iva_emisor o cliente_codigo, y no aparecen en ninguna
        # columna, así que avisar de esas es solo ruido).
        campos_vacios = [
            k for k, v in resultado["cabecera"].items()
            if v is None and k in CABECERA_COLUMNAS
        ]
        mensajes_extra = []
        if campos_vacios:
            mensajes_extra.append(
                f"no se pudieron detectar estos campos: {', '.join(campos_vacios)}"
            )
        if ruta_items_debug and not resultado["items"]:
            mensajes_extra.append("no se detectó la tabla de ítems")

        if mensajes_extra:
            mensaje = f"Factura cargada, pero {'; '.join(mensajes_extra)}."
            if ruta_texto_debug:
                mensaje += (
                    f" Podés revisar qué reconoció el OCR en "
                    f"uploads/{os.path.basename(ruta_texto_debug)} "
                    f"y el detalle de ítems/totales reconstruidos en "
                    f"uploads/{os.path.basename(ruta_items_debug)}."
                )
            flash(mensaje)
        else:
            flash(f"Factura '{archivo.filename}' procesada correctamente ({len(resultado['items'])} ítems detectados).")
    except pytesseract.TesseractNotFoundError:
        flash(
            "No se encontró el programa Tesseract OCR instalado en el sistema. "
            "Revás la sección 'OCR para fotos/imágenes' en el README para instalarlo."
        )
    except Exception as e:
        flash(f"Error procesando el archivo: {e}")

    return redirect(url_for("index"))


@app.route("/descargar")
def descargar():
    if not os.path.exists(EXCEL_PATH):
        flash("Todavía no se generó ningún Excel.")
        return redirect(url_for("index"))
    return send_file(EXCEL_PATH, as_attachment=True, download_name="facturas.xlsx")


@app.route("/descargar_ultima")
def descargar_ultima():
    if not os.path.exists(ULTIMA_PATH):
        flash("Todavía no procesaste ninguna factura en esta sesión.")
        return redirect(url_for("index"))
    return send_file(ULTIMA_PATH, as_attachment=True, download_name="ultima_factura.xlsx")


@app.route("/reiniciar", methods=["POST"])
def reiniciar():
    if os.path.exists(EXCEL_PATH):
        os.remove(EXCEL_PATH)
    flash("Excel reiniciado.")
    return redirect(url_for("index"))


if __name__ == "__main__":
    app.run(debug=True, port=5000)
